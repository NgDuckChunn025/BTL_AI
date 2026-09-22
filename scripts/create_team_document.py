"""Build the requested Word assignment using the managed document runtime.

Preset: standard_business_brief. Header: memo_masthead (no rule).
Named overrides: title 23 pt/4 pt after; table text 10 pt; quiet footer 9 pt.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "report" / "Phan_cong_cong_viec_EduPredict_AI.docx"


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"
    props = t._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:w"), "9360")
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120"); indent.set(qn("w:type"), "dxa"); props.append(indent)
    margins = OxmlElement("w:tblCellMar")
    for key, value in [("top", 80), ("bottom", 80), ("start", 120), ("end", 120)]:
        e = OxmlElement("w:" + key); e.set(qn("w:w"), str(value)); e.set(qn("w:type"), "dxa"); margins.append(e)
    props.append(margins)
    for grid, width in zip(t._tbl.tblGrid.gridCol_lst, widths):
        grid.set(qn("w:w"), str(width))
    for col, width in zip(t.columns, widths): col.width = Inches(width/1440)
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    repeat = OxmlElement("w:tblHeader"); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, values): cell.text = value
    for i, row in enumerate(t.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width/1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                fill = OxmlElement("w:shd"); fill.set(qn("w:fill"), "F2F4F7"); cell._tc.get_or_add_tcPr().append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(10); run.bold = i == 0
    return t


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    styles = {"Normal": (11, "172B4D", 0, 6), "Title": (23, "0B2545", 0, 4),
              "Subtitle": (12, "667085", 0, 12), "Heading 1": (16, "2E74B5", 16, 8),
              "Heading 2": (13, "2E74B5", 12, 6), "Heading 3": (12, "1F4D78", 8, 4)}
    for name, (size, color, before, after) in styles.items():
        s = doc.styles[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.line_spacing = 1.1
        s.paragraph_format.keep_with_next = name.startswith("Heading")
    section.header.paragraphs[0].text = "EDUPREDICT AI  |  BÀI TẬP LỚN TRÍ TUỆ NHÂN TẠO"
    section.header.paragraphs[0].style = doc.styles["Subtitle"]
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Kế hoạch phân công đề xuất  |  Trang ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.add_paragraph("PHÂN CÔNG CÔNG VIỆC", "Title")
    doc.add_paragraph("EduPredict AI – Cổng học vụ và cảnh báo sớm", "Subtitle")
    doc.add_paragraph("Mục tiêu: hoàn thiện ứng dụng Streamlit, tái lập quy trình học máy trên dữ liệu học vụ mô phỏng và chuẩn bị báo cáo bảo vệ.")
    doc.add_paragraph("Phân công dưới đây là kế hoạch đề xuất, không phải xác nhận mức đóng góp đã hoàn thành. Nhóm thống nhất lịch cụ thể theo hạn nộp của môn học.")
    doc.add_heading("Thành viên và vai trò", level=1)
    table(doc, ["Thành viên", "MSSV", "Vai trò phụ trách"], [
        ["Nguyễn Đức Chung", "20231953", "Trưởng nhóm; kiến trúc, tích hợp và UI/UX"],
        ["Đặng Phúc Đình", "20231110", "Dữ liệu, tiền xử lý và huấn luyện"],
        ["Trần Quang Huy", "20231338", "Đánh giá, kiểm thử và tài liệu"]], [3000, 1600, 4760])
    doc.add_heading("Nguyễn Đức Chung – Trưởng nhóm", level=2)
    doc.add_paragraph("Chốt yêu cầu, chia nhiệm vụ và theo dõi tiến độ. Phụ trách app.py, giao diện sáng/tối, điều hướng, form dự báo, lưu kịch bản và tích hợp model. Rà soát pull request và điều phối buổi demo.")
    doc.add_paragraph("Bàn giao: ứng dụng chạy ổn định, quy ước code, hướng dẫn khởi chạy và kịch bản demo. Nghiệm thu: đủ các trang, nút thao tác hoạt động đúng, chữ/bảng/widget đọc rõ ở cả hai theme.")
    doc.add_heading("Đặng Phúc Đình", level=2)
    doc.add_paragraph("Phụ trách schema và chất lượng hai CSV học vụ; giải thích quy tắc sinh nhãn. Xây dựng Pipeline, tách tập theo MSSV, huấn luyện Logistic Regression, Random Forest, Gradient Boosting và Dummy Baseline.")
    doc.add_paragraph("Bàn giao: dữ liệu có mô tả, mã train tái lập, model và metadata. Nghiệm thu: không đưa điểm cuối kỳ/nhãn/MSSV vào feature; không trùng sinh viên giữa các tập; chạy lại train thành công.")
    doc.add_page_break()
    doc.add_heading("Trần Quang Huy", level=2)
    doc.add_paragraph("Phụ trách Accuracy, Recall, Precision, F1, ROC-AUC, confusion matrix và phân tích lỗi. Kiểm thử form, khôi phục hồ sơ, lịch sử, xuất file; viết README, hướng dẫn chỉ số, báo cáo và slide.")
    doc.add_paragraph("Bàn giao: bảng benchmark, bộ kiểm thử, hướng dẫn sử dụng và phần đánh giá trong báo cáo. Nghiệm thu: số liệu giao diện khớp artifact; giải thích được cảnh báo nhầm/bỏ sót; báo cáo nêu rõ giới hạn dữ liệu mô phỏng.")
    doc.add_heading("Lộ trình phối hợp", level=1)
    table(doc, ["Mốc", "Đầu việc", "Chủ trì"], [
        ["Giai đoạn 1", "Chốt yêu cầu, schema, tiêu chí nghiệm thu", "Chung + cả nhóm"],
        ["Giai đoạn 2", "Kiểm tra dữ liệu, tiền xử lý, train model", "Đình"],
        ["Giai đoạn 3", "Hoàn thiện UI và tích hợp dự báo", "Chung"],
        ["Giai đoạn 4", "Kiểm thử, benchmark, hoàn thiện tài liệu", "Huy"],
        ["Trước bảo vệ", "Chạy thử trên máy khác, luyện thuyết trình", "Cả nhóm"]], [1850, 5060, 2450])
    doc.add_heading("Quy tắc bàn giao và kiểm tra chéo", level=1)
    doc.add_paragraph("Chung kiểm tra tích hợp model và giao diện; Đình kiểm tra cách diễn giải dữ liệu/chỉ số; Huy kiểm tra khả năng tái lập và các thao tác người dùng. Mỗi thay đổi cần mô tả đầu vào, đầu ra và cách kiểm tra.")
    doc.add_paragraph("Tỷ trọng kế hoạch đề xuất: Chung 34%, Đình 33%, Huy 33%. Đánh giá đóng góp thực tế dựa trên sản phẩm bàn giao, lịch sử commit và kết quả kiểm tra chéo.")
    doc.add_heading("Chuẩn bị bảo vệ", level=1)
    doc.add_paragraph("Chung: giới thiệu bài toán, kiến trúc và demo. Đình: dữ liệu, feature, pipeline và cách train. Huy: benchmark, phân tích lỗi, giới hạn và hướng phát triển. Cả ba cùng nắm luồng dữ liệu từ CSV đến dự báo.")
    doc.add_paragraph("Hướng phát triển ưu tiên: calibration xác suất và tối ưu ngưỡng trên validation; lưu lịch sử bền vững; đánh giá theo kỳ học khi có dữ liệu phù hợp. Chỉ đưa vào sản phẩm khi có tiêu chí kiểm chứng rõ ràng.")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    # Structural audit: names/IDs and explicit table geometry.
    loaded = Document(OUTPUT)
    texts = " ".join(p.text for p in loaded.paragraphs) + " ".join(c.text for t in loaded.tables for row in t.rows for c in row.cells)
    for value in ["Nguyễn Đức Chung", "20231953", "Đặng Phúc Đình", "20231110", "Trần Quang Huy", "20231338"]:
        assert value in texts
    for t in loaded.tables:
        assert sum(int(g.get(qn("w:w"))) for g in t._tbl.tblGrid.gridCol_lst) == 9360
    print(OUTPUT)


if __name__ == "__main__": build()
